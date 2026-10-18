"""Build Invoices-base.xlsx: Invoice-Data (clients + line items) and
Invoice-Template (formatted layout the macro copies per client).
100% local. by Muhammad Bilal."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Invoices-base.xlsx")

NAVY = "1F4E79"
LIGHT = "D9E2F3"
TITLE_FONT = Font(name="Calibri", size=20, bold=True, color=NAVY)
HDR_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor=NAVY)
LBL_FONT = Font(name="Calibri", size=11, bold=True, color=NAVY)
VAL_FONT = Font(name="Calibri", size=11)
TOT_FONT = Font(name="Calibri", size=12, bold=True, color=NAVY)
THIN = Border(*(Side(style="thin", color="9AA5B8") for _ in range(4)))
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

CLIENTS = [
    ("C001", "Riverdale Traders", "12 Market Road", "Lahore", 16),
    ("C002", "Harbor Supplies Co.", "45 Dock Street", "Karachi", 16),
    ("C003", "Maple Foods Ltd.", "78 Orchard Lane", "Islamabad", 16),
]

ITEMS = [
    ("C001", "Project management software license", 5, 45000),
    ("C001", "Staff training sessions", 3, 60000),
    ("C002", "Office workstation setup", 10, 85000),
    ("C002", "Network cabling (per point)", 24, 3500),
    ("C003", "Cold-chain audit consultancy", 1, 250000),
    ("C003", "Temperature loggers", 8, 12500),
]


def header_row(ws, row, cols):
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = THIN


def build_data(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "CLIENTS"
    ws["A1"].font = TITLE_FONT
    header_row(ws, 2, ["ClientID", "ClientName", "Address", "City", "TaxPct"])
    for i, cl in enumerate(CLIENTS, 3):
        for c, v in enumerate(cl, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = VAL_FONT
            cell.alignment = LEFT
            cell.border = THIN
    ws["G1"] = "LINE ITEMS"
    ws["G1"].font = TITLE_FONT
    header_row_gj(ws)
    for i, it in enumerate(ITEMS, 3):
        for c, v in enumerate(it, 7):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = VAL_FONT
            cell.alignment = LEFT if c == 8 else CENTER
            cell.border = THIN
    for col, w in zip("ABCDE", (10, 24, 22, 14, 10)):
        ws.column_dimensions[col].width = w
    for col, w in zip("GHIJ", (10, 38, 8, 14)):
        ws.column_dimensions[col].width = w


def header_row_gj(ws):
    for c, val in enumerate(["ClientID", "Description", "Qty", "Rate (PKR)"], 7):
        cell = ws.cell(row=2, column=c, value=val)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = THIN


def build_template(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 20

    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = "TAX INVOICE"
    t.font = TITLE_FONT
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    labels = [("A3", "Invoice No:"), ("A4", "Date:"), ("A5", "Bill To:")]
    for addr, txt in labels:
        c = ws[addr]
        c.value = txt
        c.font = LBL_FONT
        c.alignment = LEFT
    for addr in ("C3", "C4", "C5", "C6"):
        ws[addr].font = Font(name="Calibri", size=11, bold=True)
        ws[addr].alignment = LEFT
    ws["C3"].value = "[INVNO]"
    ws["C4"].value = "[DATE]"
    ws["C5"].value = "[CLIENT]"
    ws["C6"].value = "[ADDRESS]"

    header_row(ws, 8, ["#", "Description", "Qty", "Rate (PKR)", "Amount (PKR)"])
    for r in range(9, 15):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = CENTER if c != 2 else LEFT
            cell.font = VAL_FONT
        ws.row_dimensions[r].height = 20

    for r, txt in ((15, "Subtotal"), (16, "Tax (16%)"), (17, "GRAND TOTAL")):
        ws.cell(row=r, column=4, value=txt).font = LBL_FONT
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="right", vertical="center")
        v = ws.cell(row=r, column=5)
        v.font = TOT_FONT
        v.number_format = '#,##0'
        v.alignment = CENTER
    ws.row_dimensions[17].height = 26
    ws["A19"] = "Thank you for your business."
    ws["A19"].font = Font(name="Calibri", size=10, italic=True, color="666666")


def main():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Invoice-Data"
    build_data(ws1)
    ws2 = wb.create_sheet("Invoice-Template")
    build_template(ws2)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
